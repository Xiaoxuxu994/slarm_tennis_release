"""Build an editable project-report deck; no model/data dependencies."""

from pathlib import Path
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.xmlchemy import OxmlElement
from pptx.oxml.ns import qn
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs/TENNIS_PROJECT_REPORT_codex.pptx"
BG = "F5F7F8"
INK = "20282D"
MUTED = "596970"
TEAL = "087F78"
CORAL = "C75849"
BLUE = "316DA0"
PALE = "E4EFEE"
WHITE = "FFFFFF"
FONT = "Microsoft YaHei"
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)


def rect(slide, x, y, w, h, color, line=None):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(color)
    shape.line.fill.background()
    if line:
        shape.line.color.rgb = RGBColor.from_string(line)
    return shape


def text(slide, x, y, w, h, value, size=20, color=INK, bold=False):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.02)
    tf.margin_top = tf.margin_bottom = 0
    for i, line in enumerate(value.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(8)
        p.font.name = FONT
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = RGBColor.from_string(color)
        p.text = line
        # Explicit East Asian face avoids theme substitution on Chinese text.
        for run in p.runs:
            rpr = run._r.get_or_add_rPr()
            ea = OxmlElement("a:ea")
            ea.set("typeface", FONT)
            rpr.append(ea)
    return shape


def page(kicker, title, subtitle, source, notes):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor.from_string(BG)
    rect(slide, 0, 0, 0.15, 7.5, TEAL)
    text(slide, 0.55, 0.28, 12, 0.3, kicker, 11, TEAL, True)
    text(slide, 0.55, 0.82, 12.2, 0.68, title, 30, INK, True)
    text(slide, 0.55, 1.61, 12.1, 0.55, subtitle, 16, MUTED)
    rect(slide, 0.55, 6.92, 12.1, 0.015, "CBD7DA")
    text(slide, 0.55, 7.06, 11.55, 0.26, source, 9, MUTED)
    text(slide, 12.18, 7.02, 0.65, 0.3, f"{len(prs.slides):02d}", 13, TEAL, True)
    slide.notes_slide.notes_text_frame.text = notes + "\n来源：" + source
    return slide


def block(slide, x, y, w, title, body, accent=TEAL, h=1.4):
    rect(slide, x, y, w, h, WHITE)
    rect(slide, x, y, 0.04, h, accent)
    text(slide, x + 0.18, y + 0.14, w - 0.36, 0.4, title, 19, accent, True)
    text(slide, x + 0.18, y + 0.65, w - 0.36, h - 0.68, body, 16)


def arrow(slide, x, y, w=0.45, color=TEAL):
    shape = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(y), Inches(w), Inches(0.27))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(color)
    shape.line.fill.background()


def banner(slide, message, color=TEAL):
    rect(slide, 0.55, 6.15, 12.1, 0.52, color)
    text(slide, 0.75, 6.25, 11.7, 0.32, message, 16, WHITE, True)


def table(slide, headers, rows, widths, y=2.35, row_h=0.63):
    x = 0.55
    for value, width in zip(headers, widths):
        rect(slide, x, y, width, row_h, TEAL)
        text(slide, x + 0.12, y + 0.13, width - 0.24, row_h - 0.14, value, 15, WHITE, True)
        x += width
    for j, row in enumerate(rows):
        x = 0.55
        for value, width in zip(row, widths):
            yy = y + (j + 1) * row_h
            rect(slide, x, yy, width, row_h, WHITE if j % 2 == 0 else "EAF0F2")
            text(slide, x + 0.12, yy + 0.13, width - 0.24, row_h - 0.14, value, 15)
            x += width


def build():
    s = page("SLARM / TENNIS CATCHING / PROJECT REPORT", "从动态重建到网球接球",
             "三视角因果感知 · 物理运动外推 · 对象级 latent 接口", "2026-09-17 | 本地代码 4f7e46a + 工作区 | 基于 kevinchiu19/SLARM",
             "项目不是重新发明SLARM，而是围绕接球的时间约束、小目标几何和任务读出进行系统性改造。当前主线是pixel/MS3；token分支作为探索。")
    for x, title, body, color in [
        (0.55, "01  看清球", "三视角几何\n球区域专用监督", TEAL),
        (4.68, "02  预测未来", "终端高斯外推\n物理单位与坐标统一", BLUE),
        (8.81, "03  输出可用表征", "显式球状态 + 对象 latent\n面向下游策略接口", CORAL),
    ]:
        block(s, x, 2.7, 3.84, title, body, color, 2.15)
    banner(s, "核心定位：在开源世界模型之上，构建可验证的接球感知与预测链路")

    s = page("01 / CONTRIBUTION BOUNDARY", "继承基础能力，新增任务闭环",
             "将项目贡献与上游能力分开，避免把工程实现表述为未经证实的学术首创。",
             "上游：https://github.com/kevinchiu19/SLARM | 本地 src/models/slarm.py、src/utils/stream25_losses.py",
             "上游已有动态GS、场景运动、流式与语言对齐能力。本项目的重点是终端外推归属、小球监督、task semantic和对象latent接口。未做上游所有版本逐行审计。")
    table(s, ["上游基础", "本项目改造", "证据状态"], [
        ["动态 GS / 运动预测", "终端动态高斯承担未来外推", "已实现 / 启用"],
        ["重建与语义学习", "球区域几何 + 物理监督；关闭LSeg监督", "已实现 / 启用"],
        ["流式多视角 latent", "in-trunk ball token + 状态/特征接口", "已实现 / 探索"],
        ["通用重建指标", "固定接球时刻误差 + 多因素消融", "已实现"],
    ], [3.0, 6.5, 2.6])
    banner(s, "贡献是任务化机制与验证体系，不是重新提出 GS 或流式 Transformer")

    s = page("02 / SYSTEM", "两条读出路径，共享因果视觉主干",
             "当前 10k 主线使用 pixel/MS3；ball-token 是可选分支，不是默认同时启用。",
             "src/models/slarm.py | src/models/stream_session.py | docs/BALL_LATENT_EXPORT.md",
             "输入为六个时刻三视角RGB和标定。像素路径经过渲染、语义选球、深度反投影和状态聚合。token路径提供状态和三目1536维特征。下游DynamicVLA尚未实现验证闭环。")
    block(s, 0.55, 2.65, 2.5, "三视角 × 六时刻", "RGB + 相机标定\n0/3/6/9/12/15", h=1.7)
    arrow(s, 3.13, 3.35)
    block(s, 3.7, 2.65, 3.0, "SLARM 流式主干", "patch + 几何/时间编码\n多视角与历史聚合", BLUE, 1.7)
    arrow(s, 6.8, 3.35)
    block(s, 7.45, 2.35, 5.2, "主线：Pixel / Gaussian / MS3", "未来渲染 → 选球反投影 → 状态 → 接球点", TEAL, 1.4)
    block(s, 7.45, 4.03, 5.2, "探索：Ball Token", "3 × 1536 latent → p / v；策略接口待验证", CORAL, 1.4)
    banner(s, "同时保留可解释几何读出与紧凑对象表征，但分别评估其价值")

    s = page("03 / CAUSAL EXTRAPOLATION", "让最后一次观测承担未来预测",
             "terminal_context_extrapolation：明确终端及后续目标的高斯时间归属。",
             "src/models/slarm.py::forward_renderer | src/models/temporal_ownership.py",
             "终端及未来清除更早高斯的时间归属重叠，终端动态高斯按MS3推进。0.3秒和1秒为默认30fps窗口。frame45参考是解析弹道，不是保存的图像标注。")
    for x, title, body, color in [(0.55, "因果观测", "frame0 → frame15\n只使用已经到达的图像", TEAL),
                                  (4.68, "短时预测", "frame15 → frame24\n0.3 s；有图像/轨迹标注", BLUE),
                                  (8.81, "接球时刻", "frame15 → frame45\n1.0 s；解析状态参考", CORAL)]:
        block(s, x, 2.5, 3.84, title, body, color, 1.65)
    text(s, 0.8, 4.65, 11.7, 0.65, "p(t+dt) = p(t) + v dt + 0.5 a dt^2 + (1/6) j dt^3", 25, TEAL, True)
    banner(s, "关键改造：从已观测时刻重建，走向受因果约束的任务时刻预测")

    s = page("04 / SMALL-OBJECT SUPERVISION", "全图重建之外，单独照顾网球",
             "背景占据绝大多数像素；全图 RGB 好看，并不保证小球几何与速度准确。",
             "src/utils/stream25_losses.py | src/dataset/stream25.py::build_dense_ms3_gt",
             "球区域RGB、米制深度及尾部误差加强小球监督，球/静态MS3分开。多视角通过共享表征和渲染监督隐式约束，目前不宣称存在显式两两三视角一致性loss。")
    block(s, 0.55, 2.45, 3.84, "全局重建", "RGB + LPIPS\n相对深度 / 语义 / opacity", TEAL, 2.0)
    block(s, 4.68, 2.45, 3.84, "球局部几何", "球区域 RGB\n米制深度：均值 + 误差尾部", BLUE, 2.0)
    block(s, 8.81, 2.45, 3.84, "运动与静态分离", "球：v / a / j 监督\n静态：独立采样与零运动约束", CORAL, 2.0)
    text(s, 0.75, 5.0, 11.8, 0.6, "三视角共享表示与渲染提供隐式约束；不是新增的显式两两一致性 loss。", 18, MUTED)
    banner(s, "设计目标：避免大面积背景损失掩盖少像素网球的误差")

    s = page("05 / PHYSICS CONTRACT", "物理先验分三层，不能混为一谈",
             "统一物理单位、坐标变换和参考时刻，比单纯增大一个 loss 更重要。",
             "stream25_losses.py | slarm.py::_apply_physics_ms3_override | stream25_metrics.py",
             "当前10k配置没有显式开启ms3_physics_override，默认关闭。物理监督与可选硬覆盖要分清。token速度scale仅为loss归一化，不是decoder缩放。learned MS3已接近重力，不应夸大硬覆盖收益。")
    table(s, ["层次", "实现", "当前状态"], [
        ["训练监督", "球 MS3 GT：速度 / 重力 / 零 jerk", "pixel 主线启用"],
        ["解析状态推进", "ball-token 轨迹 / catch 使用已知重力", "token 分支按配置启用"],
        ["可选硬约束", "动态 a=g、j=0 的 physics override", "10k 配置未显式启用"],
        ["单位契约", "状态 m、m/s；时间秒；canonical → rig", "全链路检查"],
    ], [2.1, 6.8, 3.2])
    banner(s, "物理模型减少不必要自由度；精度仍依赖位置与速度的感知泛化")

    s = page("06 / TASK SEMANTICS", "关闭 LSeg 特征监督，保留任务语义",
             "从通用语言对齐转为封闭接球场景所需的语义输出。",
             "configs/exp0915_001_slarm_stream25_0908_10k_pixel_finetune.yml | slarm.py::forward_task_semantic_predictor",
             "配置online_feat=false，enable_feat_loss=false，lseg权重0，四类task semantic开启。仓库保留兼容逻辑，不是删除一切LSeg代码。没有速度/显存统一对照，不能写加速倍数。")
    block(s, 0.55, 2.5, 5.65, "关闭的部分", "在线 LSeg 特征监督\n语言特征蒸馏损失", CORAL, 2.1)
    block(s, 7.0, 2.5, 5.65, "保留 / 新增的部分", "四类 task semantic + 渲染监督\n预测球语义 → 像素路径选球", TEAL, 2.1)
    arrow(s, 6.35, 3.35)
    text(s, 0.75, 5.03, 11.6, 0.6, "取舍：减少任务不需要的监督依赖；不再宣称原样保留开放词汇能力。", 19, MUTED)
    banner(s, "这是任务化简化；目前不报告未经测量的推理加速或显存下降比例")

    s = page("07 / OBJECT-CENTRIC REPRESENTATION", "Ball Token：从像素场到对象状态",
             "in-trunk token 参与主干 attention，提供显式 p/v 与三目 latent 两类接口。",
             "aggregator.py | slarm.py | docs/BALL_LATENT_EXPORT.md | stream25_losses.py",
             "terminal三目token已经交换信息，不是独立单目估计。pooled基线mean后共享MLP输出p/v。位置、速度、轨迹、landing监督对齐接球目标。没有证据证明当前token优于pixel或latent下游质量已达标。")
    block(s, 0.55, 2.5, 3.35, "In-trunk token", "与 image patches 联合聚合\n终端输出 [B,3,1536]", TEAL, 1.75)
    arrow(s, 4.02, 3.2)
    block(s, 4.6, 2.5, 3.35, "状态读出", "三目 mean → 共享 MLP\np15、v15（物理单位）", BLUE, 1.75)
    arrow(s, 8.07, 3.2)
    block(s, 8.65, 2.5, 4.0, "任务监督", "位置 + 速度\n轨迹 + 固定时刻接球点", CORAL, 1.75)
    text(s, 0.75, 4.88, 11.7, 0.8, "三目 latent 可单独导出 → 下游 adapter / DynamicVLA\n接口已具备；融合位置、策略收益与特征质量尚待验证。", 20)
    banner(s, "已实现的表征接口，不等于已证明优于像素路径的精度方案", CORAL)

    s = page("08 / CONTROLLED MOTION PROBE", "冻结基线，只学习历史速度修正",
             "006–010：从原 ball-token baseline007999 初始化，不叠加004 temporal-joint head。",
             "src/models/ball_velocity_residual.py | configs/exp0911_010_balltoken_history_concat_deltat_diff_regularized.yml",
             "valid-view mean并非GT可见性Oracle；无外部mask时依赖有限值。共享投影及LN后计算差分，末层零初始化，物理速度相加。只有新增head训练，prefix关闭。结果暂不构成精度突破。")
    block(s, 0.55, 2.45, 5.8, "六帧历史输入", "valid-view mean → 共享1536→256 → LN\n[h_t, h_terminal−h_t, Δt秒, valid]", TEAL, 1.65)
    block(s, 6.85, 2.45, 5.8, "低自由度残差读出", "concat 3084 → 512 → 256 → 3\nv_final = stopgrad(v_base) + Δv", BLUE, 1.65)
    text(s, 0.75, 4.5, 11.8, 0.9, "零初始化  /  baseline冻结  /  不加prefix监督  /  轻微Δv正则\n010验证速度约0.101/0.268 m/s，与baseline接近，尚无稳定收益。", 20)
    banner(s, "价值：可控地验证历史 latent 的可读性；结论：暂不作为主模型", CORAL)

    s = page("09 / EVIDENCE", "当前候选：接球时刻中位误差 7.0 cm",
             "最新训练：019999（待评测）；下表仍为013999 / offset=0，不混用权重与成绩。",
             "用户提供的 ckpt13999 aggregate | manifest / 样本数及同场景对照待补全",
             "最新训练checkpoint为019999，下表已知成绩来自013999。不能从不同数据报告直接计算微调收益百分比。表面补偿关闭，catch为固定时刻解析参考。用户另报013999成功率67.5%、n=200，判据待核对，不直接列为机器人成功率。同场景checkpoint应采用配对统计，没有必须差10个百分点的通用标准。")
    table(s, ["指标", "Median", "P95", "解释"], [
        ["frame24 position", "4.1 cm", "11.2 cm", "短时预测，历史表面口径"],
        ["catch position", "7.0 cm", "17.4 cm", "固定接球时刻，约1秒外推"],
        ["ball MS3 velocity", "0.067 m/s", "0.185 m/s", "球区域运动场指标"],
        ["ball depth farthest", "2.2 cm", "3.8 cm", "球区域深度误差"],
    ], [3.3, 2.0, 2.0, 4.8])
    banner(s, "可以报告候选精度；不能将其直接换算为机器人成功率或同分布提升比例")

    s = page("10 / ABLATIONS", "负结果让下一步更明确",
             "不是模块越多越好：优先保留已验证的主路径，停止无收益叠加。",
             "用户提供004读出/历史拟合结果、010训练与验证报告 | docs/EXPERIMENTS_AND_ERROR_BUDGET.md",
             "旧文档的几何极限推断已被train/val差距推翻；免费滑窗推断也被offset实验挑战。004fit比较同一checkpoint。010无明显收益，不代表latent一定没有信息。")
    table(s, ["验证方向", "观察", "决策"], [
        ["逐目位置再平均", "多数场景没有胜过 feature mean", "保留原读出"],
        ["004历史位置拟合", "速度中位0.1074 → 0.1465 m/s", "不替换速度头"],
        ["冻结历史残差010", "修正小，验证速度接近baseline", "暂不升级主模型"],
        ["直接平移窗口+3", "catch中位0.665 m；配对待确认", "排查时间 / 窗口泛化"],
    ], [3.1, 6.0, 3.0])
    banner(s, "不再宣称已经达到几何极限，也不宣称后移观测能免费提升精度", CORAL)

    s = page("11 / EVALUATION & VISUALIZATION", "从好看的重建，走向可追溯的任务评价",
             "误差分解与可视化各司其职：前者决定结论，后者帮助解释。",
             "tools/verify_physics_extrapolation.py | scripts/visualize_ball_tokens.py | tools/export_gaussian_sequence.py",
             "架构图为可编辑形状，未使用伪造推理图。GS可导PLY查看几何，token可导attention投影。attention不等于因果解释。演示最终建议替换为同checkpoint、同scene、同frame的真实输出。")
    block(s, 0.55, 2.45, 3.84, "误差诊断", "pixel / token\nfree / gravity / linear\n视线 / 横向 / 逐轴", TEAL, 2.4)
    block(s, 4.68, 2.45, 3.84, "GS 可视化", "动态渲染视频\n逐帧 Gaussian PLY\n球预测轨迹与 GT 对照", BLUE, 2.4)
    block(s, 8.81, 2.45, 3.84, "Token 可视化", "attention 投影到图像\n状态轨迹与速度残差\nlatent 导出与探针评估", CORAL, 2.4)
    banner(s, "评估需标记：缺样本、球面补偿、预测时间、坐标系与 checkpoint 来源")

    s = page("12 / NEXT STEPS", "下一步：先解决泛化，再验证策略价值",
             "不把新增模块当作进步；每一步都对应可检验的问题。",
             "docs/RANDOM_WINDOW_FINETUNE_codex.md（方案未实现） | docs/BALL_LATENT_EXPORT.md",
             "固定同场景比较旧模型和13999。随机offset0至3，保持timespan0.8，同步target及GT。25帧限制未来监督，offset3只剩0.2秒；要0.3秒需数据至少到27。DynamicVLA尚无已验证融合结果。")
    for y, number, title, body in [
        (2.35, "01", "补齐同场景对照", "旧模型 vs 13999；记录 catch p95、距离 hit rate 与延迟"),
        (3.45, "02", "验证窗口泛化", "同帧MS3检查 → 随机offset0/1/2/3；保留offset0表现"),
        (4.55, "03", "验证 latent 与控制价值", "冻结特征探针 → DynamicVLA adapter → 机器人闭环"),
    ]:
        text(s, 0.65, y, 0.65, 0.5, number, 28, TEAL, True)
        text(s, 1.5, y, 10.9, 0.36, title, 20, INK, True)
        text(s, 1.5, y + 0.43, 10.9, 0.42, body, 17, MUTED)
    banner(s, "项目产出：可预测、可解释、可验证的接球感知基础，而非已完成的策略闭环")

    s = page("APPENDIX / REPRODUCIBILITY", "代码、配置与证据索引",
             "详细边界、指标表与一分钟汇报稿见配套 Markdown；每页附有讲稿备注。",
             "docs/TENNIS_PROJECT_REPORT_codex.md | tools/build_tennis_project_report.py",
             "所有数值来自项目文档及用户贴出的报告。配套MD列出完整文件路径和实验边界。上游归属SLARM。没有本地GPU实测、机器人闭环、下游action收益或统一性能benchmark。")
    table(s, ["内容", "定位"], [
        ["开源基础", "https://github.com/kevinchiu19/SLARM"],
        ["当前 pixel 配置", "exp0915_001_slarm_stream25_0908_10k_pixel_finetune.yml"],
        ["对象基线 / 残差", "exp0908_003 / exp0911_010 对应 configs"],
        ["核心实现", "slarm.py / stream25_losses.py / ball_velocity_residual.py"],
        ["实验与接口说明", "EXPERIMENTS_AND_ERROR_BUDGET.md / BALL_LATENT_EXPORT.md"],
    ], [3.1, 9.0], row_h=0.58)
    banner(s, "可用于最终汇报；有条件的结果与未完成的验证均已明确标记")


def validate():
    """Check packaging, slide count, notes and geometric bounds."""
    deck = Presentation(OUT)
    assert len(deck.slides) == 14
    notes_ids = deck._element.find(qn("p:notesMasterIdLst"))
    assert notes_ids is not None and len(notes_ids) == 1
    assert deck.part.rels[notes_ids[0].get(qn("r:id"))].reltype == RT.NOTES_MASTER
    for slide in deck.slides:
        assert slide.notes_slide.notes_text_frame.text
        for shape in slide.shapes:
            assert shape.left >= 0 and shape.top >= 0
            assert shape.left + shape.width <= deck.slide_width + 2
            assert shape.top + shape.height <= deck.slide_height + 2
    print(f"Validated {len(deck.slides)} slides: {OUT}")


def register_notes_master():
    """python-pptx creates a notes relationship but omits its presentation entry."""
    root = prs._element
    if root.find(qn("p:notesMasterIdLst")) is None:
        relation = next(r for r in prs.part.rels.values() if r.reltype == RT.NOTES_MASTER)
        ids = OxmlElement("p:notesMasterIdLst")
        item = OxmlElement("p:notesMasterId")
        item.set(qn("r:id"), relation.rId)
        ids.append(item)
        root.insert(list(root).index(root.find(qn("p:sldSz"))), ids)
    root.find(qn("p:sldSz")).set("type", "custom")


def preview():
    """Approximate layout proof, not a PowerPoint/Keynote rendering."""
    from PIL import Image, ImageDraw, ImageFont

    font_path = Path("/System/Library/Fonts/Hiragino Sans GB.ttc")
    if not font_path.exists():
        print("Preview skipped: install a CJK font and update font_path")
        return
    deck = Presentation(OUT)
    scale = 100 / 914400
    tiles = []
    warnings = []
    for index, slide in enumerate(deck.slides, 1):
        img = Image.new("RGB", (1334, 750), "#" + BG)
        draw = ImageDraw.Draw(img)
        for shape in slide.shapes:
            x, y, w, h = [round(v * scale) for v in
                           (shape.left, shape.top, shape.width, shape.height)]
            if shape.fill.type == 1:
                draw.rectangle((x, y, x + w, y + h), fill="#" + str(shape.fill.fore_color.rgb))
            if not shape.has_text_frame or not shape.text:
                continue
            tf = shape.text_frame
            yy = y + round(tf.margin_top * scale)
            xx = x + round(tf.margin_left * scale)
            avail = w - round((tf.margin_left + tf.margin_right) * scale)
            for para in tf.paragraphs:
                size = round((para.font.size.pt if para.font.size else 20) * 100 / 72)
                font = ImageFont.truetype(str(font_path), size)
                line = ""
                lines = []
                for char in para.text:
                    if line and draw.textlength(line + char, font=font) > avail:
                        lines.append(line)
                        line = ""
                    line += char
                lines.append(line)
                color = "#" + str(para.font.color.rgb)
                for line in lines:
                    draw.text((xx, yy), line, font=font, fill=color)
                    yy += round(size * 1.16)
                yy += round((para.space_after.pt if para.space_after else 0) * 100 / 72)
            if yy - 11 > y + h + 4:
                warnings.append((index, shape.text[:45]))
        img.save(Path("/private/tmp") / f"tennis_report_slide_{index:02d}.png")
        tiles.append(img.resize((667, 375)))
    sheet = Image.new("RGB", (1334, 375 * 7), "white")
    for i, img in enumerate(tiles):
        sheet.paste(img, ((i % 2) * 667, (i // 2) * 375))
    sheet.save("/private/tmp/tennis_report_contact_sheet.png")
    print("Approximate CJK text-fit warnings:", warnings)


if __name__ == "__main__":
    build()
    register_notes_master()
    prs.core_properties.title = "从动态重建到网球接球：SLARM项目汇报"
    prs.core_properties.subject = "Task-specific algorithm adaptations, evidence and limitations"
    prs.core_properties.author = "SLARM Tennis Project"
    prs.save(OUT)
    validate()
    if "--preview" in sys.argv:
        preview()
